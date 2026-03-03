import openpyxl

wb = openpyxl.load_workbook(r'f:\for PORTFOLIO\Project_PokemonV3\Updated_PokeTables.xlsx')
ws = wb['Consolidated (V1 & V2)']

print("N2 value:", ws['N2'].value)
print("R2 value:", ws['R2'].value)
print("I2 value:", ws['I2'].value)
print("K2 value:", ws['K2'].value)
